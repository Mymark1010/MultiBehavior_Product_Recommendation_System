"""Export the data-only report to an editable Word document."""
from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/BAO_CAO_XU_LY_DU_LIEU_VI.md"


def inline(paragraph, text, size=None):
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        bold = token.startswith("**") and token.endswith("**")
        code = token.startswith("`") and token.endswith("`")
        run = paragraph.add_run(token[2:-2] if bold else token[1:-1] if code else token)
        run.font.name = "Times New Roman"
        if size is not None:
            run.font.size = Pt(size)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(10)


def build(source=SOURCE):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.widow_control = True
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("194F59")
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(10)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith("|"):
            rows = []
            while True:
                cells = [c.strip() for c in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                if index >= len(lines) or not lines[index].strip().startswith("|"):
                    break
                line = lines[index].strip()
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for number, values in enumerate(rows):
                row = table.add_row()
                if number == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for cell, value in zip(row.cells, values):
                    p = cell.paragraphs[0]
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(3)
                    inline(p, value, 10.5)
                    if number == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E7F0F2")
                        cell._tc.get_or_add_tcPr().append(shade)
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        picture = re.fullmatch(r"!\[(.*)\]\(([^)]+)\)", line)
        if picture:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((source.parent / picture[2]).resolve()), width=Cm(16.5))
            doc.add_paragraph(picture[1], style="Caption")
            continue
        heading = re.match(r"^(#{1,3}) (.*)$", line)
        if heading:
            style = {1: "Title", 2: "Heading 1", 3: "Heading 2"}[len(heading[1])]
            p = doc.add_paragraph(style=style)
            inline(p, heading[2])
        else:
            inline(doc.add_paragraph(), line)
    doc.core_properties.title = "RetailRocket — next-step theo phiên"
    doc.core_properties.author = "Vũ Hải Minh"
    doc.core_properties.subject = "Dữ liệu và phương pháp cho gợi ý tương tác kế tiếp"
    output = source.with_suffix(".docx")
    doc.save(output)
    check = Document(output)
    assert len(check.paragraphs) > 0
    print(f"Created {output.name}; {len(check.tables)} tables.")


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=SOURCE)
    build(parser.parse_args().source.resolve())


if __name__ == "__main__":
    main()
